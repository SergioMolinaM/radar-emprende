"""Informe mensual del Registro de Empresas y Sociedades (Ministerio de Economía): lo que el CSV no trae.

Del archivo «figuras y cuadros» de cada informe se toman tres series mensuales desde mayo de 2013:
- Figura_2: sociedades constituidas por escritura publicada en el Diario Oficial (fuera del Registro).
- Figura_6: socios de las sociedades constituidas en el Registro, por sexo.
- Figura_7 y Figura_9: sociedades según el sexo y la nacionalidad de sus socios.

El informe se busca en la página de la categoría (el nombre de los archivos cambia de mes a mes).
Control positivo: la columna del Registro de la Figura_2 tiene que reproducir la serie mensual de
datos/constituciones_comuna.json (scripts/constituciones.py); si difiere más de lo tolerado, se detiene.

Salida: datos/economia_res.json. Uso: python scripts/economia_res.py
"""
import datetime
import json
import pathlib
import re
import urllib.request

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
CRUDO = RAIZ / 'data-raw' / 'economia'
SALIDA = RAIZ / 'datos' / 'economia_res.json'
CONSTITUCIONES = RAIZ / 'datos' / 'constituciones_comuna.json'
CATEGORIA = 'https://www.economia.gob.cl/category/estudios-encuestas/registro-de-empresas-y-sociedades'
UA = {'User-Agent': 'Mozilla/5.0 (Radar Emprende)'}
MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
# Tolerancia del control: el CSV y el informe salen de extracciones distintas del mismo registro.
# El 5-oct-2026, 150 de 160 meses eran idénticos y el resto difería en 1 a 3 sociedades.
MAX_DIF_MES = 5
MIN_IGUALES = 0.9


def leer(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120) as r:
        return r.read()


def informe_mas_reciente():
    """(url de la nota, url del xlsx, año, mes) del informe más reciente listado en la categoría."""
    html = leer(CATEGORIA).decode('utf-8', 'replace')
    notas = set(re.findall(r'https://www\.economia\.gob\.cl/\d{4}/\d{2}/\d{2}/informe-de-creacion-de-empresas-y-cooperativas-([a-z]+)-(\d{4})\.htm', html))
    if not notas:
        raise SystemExit(f'No se encontraron informes en {CATEGORIA}: revisar la página.')
    mes, anio = max(notas, key=lambda t: (int(t[1]), MESES.index(t[0])))
    nota = re.search(rf'https://www\.economia\.gob\.cl/\d{{4}}/\d{{2}}/\d{{2}}/informe-de-creacion-de-empresas-y-cooperativas-{mes}-{anio}\.htm', html).group(0)
    xlsx = re.findall(r'https?://[^"\']+figuras-y-cuadros[^"\']+\.xlsx', leer(nota).decode('utf-8', 'replace'))
    if len(set(xlsx)) != 1:
        raise SystemExit(f'La nota {nota} no trae exactamente una planilla de figuras y cuadros: {sorted(set(xlsx))}')
    # El sitio enlaza por http; el mismo archivo responde por https (comprobado el 5-oct-2026).
    return nota, xlsx[0].replace('http://', 'https://', 1), int(anio), MESES.index(mes) + 1


def serie(ws, columnas):
    """{'AAAA-MM': [valores]} de las filas con fecha en la columna A."""
    d = {}
    for fila in ws.iter_rows(values_only=True):
        if isinstance(fila[0], datetime.datetime):
            vals = [fila[c] for c in columnas]
            if any(v is None for v in vals):
                raise SystemExit(f'{ws.title}: fila {fila[0]:%Y-%m} con celdas vacías.')
            d[fila[0].strftime('%Y-%m')] = [int(v) for v in vals]
    return d


def cabecera(ws, fila, esperado):
    vals = [ws.cell(row=fila, column=i + 2).value for i in range(len(esperado))]
    if [str(v).strip() for v in vals] != esperado:
        raise SystemExit(f'{ws.title}: la cabecera cambió. Esperado {esperado}, leído {vals}.')


def main():
    nota, url, anio, mes = informe_mas_reciente()
    if SALIDA.exists() and '--forzar' not in __import__('sys').argv:
        previo = json.loads(SALIDA.read_text(encoding='utf-8'))
        if previo.get('_meta', {}).get('archivo_url') == url:
            print(f'Sin informe nuevo: los datos ya usan {url}.')
            return
    CRUDO.mkdir(parents=True, exist_ok=True)
    ruta = CRUDO / url.rsplit('/', 1)[1]
    if not ruta.exists():
        print('bajando', url)
        ruta.write_bytes(leer(url))
    wb = openpyxl.load_workbook(ruta, data_only=True)
    actualizado = re.search(r'(\d{2})-(\d{2})-(\d{4})', str(wb['Indice']['A2'].value))
    if not actualizado:
        raise SystemExit(f'Sin fecha de actualización en Indice!A2: {wb["Indice"]["A2"].value!r}')
    actualizado = f'{actualizado.group(3)}-{actualizado.group(2)}-{actualizado.group(1)}'
    actualizado_txt = f'{actualizado[8:]}-{MESES[int(actualizado[5:7]) - 1][:3]}-{actualizado[:4]}'

    cabecera(wb['Figura_2'], 1, ['DO', 'RES', 'Total'])
    cabecera(wb['Figura_6'], 2, ['Total', 'Hombres', 'Mujeres', '(en blanco)'])
    cabecera(wb['Figura_7'], 2, ['Solo Mujeres', 'Solo Hombres', 'Hombres y Mujeres', 'No Reporta'])
    cabecera(wb['Figura_9'], 2, ['Solo Chilenas', 'Solo Extranjeras', 'Chilenas y extranjeras', 'Total'])
    do_res = serie(wb['Figura_2'], (1, 2, 3))
    socios = serie(wb['Figura_6'], (1, 2, 3, 4))
    sexo = serie(wb['Figura_7'], (1, 2, 3, 4))
    nac = serie(wb['Figura_9'], (1, 2, 3, 4))

    ultimo = f'{anio}-{mes:02d}'
    for nombre, s in [('Figura_2', do_res), ('Figura_6', socios), ('Figura_7', sexo), ('Figura_9', nac)]:
        if max(s) != ultimo:
            raise SystemExit(f'{nombre} termina en {max(s)} y el informe es de {ultimo}.')
    for k, (do, res, tot) in do_res.items():
        if do + res != tot:
            raise SystemExit(f'Figura_2 {k}: {do} + {res} != {tot}.')
    for k, (tot, h, m, b) in socios.items():
        if h + m + b != tot:
            raise SystemExit(f'Figura_6 {k}: la suma por sexo no da el total.')
    for k, (ch, ex, mix, tot) in nac.items():
        if ch + ex + mix != tot:
            raise SystemExit(f'Figura_9 {k}: la suma por nacionalidad no da el total.')

    # Control positivo contra el CSV del Registro (mismo registro, otra extracción).
    propio = json.loads(CONSTITUCIONES.read_text(encoding='utf-8'))['mensual']
    comunes = sorted(set(propio) & set(do_res))
    difs = {k: do_res[k][1] - propio[k] for k in comunes if do_res[k][1] != propio[k]}
    iguales = 1 - len(difs) / len(comunes)
    if not comunes or iguales < MIN_IGUALES or any(abs(v) > MAX_DIF_MES for v in difs.values()):
        raise SystemExit(f'La serie del Registro del informe no calza con la propia: {len(difs)} de {len(comunes)} meses distintos {difs}.')

    salida = {
        '_meta': {
            'descripcion': 'Constituciones por escritura en el Diario Oficial y composición de los socios de las sociedades constituidas en el Registro de Empresas y Sociedades, por mes.',
            'fuente': f'Ministerio de Economía, Fomento y Turismo, informe mensual de creación de empresas y cooperativas de {MESES[mes - 1]} de {anio}, planilla de figuras y cuadros (actualizada el {actualizado_txt}). Elaborada por la División de Política Comercial e Industrial con datos de la División de Empresas de Menor Tamaño.',
            'fuente_url': nota,
            'archivo_url': url,
            'metodo': 'Se copian sin cambios las series mensuales de las figuras 2 (Diario Oficial), 6 (socios por sexo), 7 (sociedades según el sexo de sus socios) y 9 (según su nacionalidad). Las sumas por año y los porcentajes son cálculo propio.',
            'nivel': 1,
            'control': f'La columna del Registro de la figura 2 reproduce la serie mensual calculada desde el CSV del Registro en {len(comunes) - len(difs)} de {len(comunes)} meses; los demás difieren en {min(map(abs, difs.values()), default=0)} a {max(map(abs, difs.values()), default=0)} sociedades (extracciones distintas del mismo registro).',
            'limites': [
                'El Diario Oficial cuenta sociedades constituidas por escritura pública; el informe no las desagrega por comuna ni por región.',
                'Los socios por sexo y nacionalidad son solo los de las sociedades del Registro, no los de las constituidas en el Diario Oficial.',
                'El informe no declara licencia de uso; se publican las series con cita de la fuente.',
            ],
            'generado': datetime.date.today().isoformat(),
            'script': 'scripts/economia_res.py',
            'licencia': 'Series del Ministerio de Economía sin licencia declarada, citadas como fuente. Sumas y porcentajes: CC BY 4.0 (Radar Emprende, de Tercera Letra SpA).',
        },
        'corte': {'anio': anio, 'mes': mes, 'actualizado': actualizado},
        'diario_oficial': {k: v[0] for k, v in do_res.items()},
        'socios_sexo': {k: dict(total=v[0], hombres=v[1], mujeres=v[2], sin_dato=v[3]) for k, v in socios.items()},
        'sociedades_sexo': {k: dict(solo_mujeres=v[0], solo_hombres=v[1], mixtas=v[2], sin_dato=v[3]) for k, v in sexo.items()},
        'sociedades_nacionalidad': {k: dict(solo_chilenas=v[0], solo_extranjeras=v[1], mixtas=v[2], total=v[3]) for k, v in nac.items()},
    }
    SALIDA.write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding='utf-8')
    print('escrito', SALIDA, '| informe', ultimo, '| control:', salida['_meta']['control'])


if __name__ == '__main__':
    main()
